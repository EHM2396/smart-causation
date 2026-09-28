"""
Router: /formulario-300 – clasificación tributaria de las operaciones a tarifa
0%, previo al reporte del Formulario 300.

Por ahora es SOLO para el administrador de la cuenta (Dayana, en el caso de
Salamanca): es la pantalla donde se toman decisiones de clasificación
tributaria compartidas por toda la firma, y la restricción evita que un
causador clasifique algo sin que quien lleva impuestos lo haya visto. Abrirlo
a causadores más adelante es cambiar una dependencia, no rehacer el módulo.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from api.dependencies import get_current_user
from db.models.auth import CuentaCliente, Empresa, Usuario
from db.models.tributario import IVA_DESCONTABLE_ESTADOS, TIPOS_ITEM, TRATAMIENTOS
from db.session import get_db
from services import formulario_300_service as f300

router = APIRouter(prefix="/formulario-300", tags=["Formulario 300"])

DB = Annotated[Session, Depends(get_db)]


def require_org_admin(current_user: Usuario = Depends(get_current_user)) -> Usuario:
    if current_user.rol not in ("org_admin", "admin"):
        raise HTTPException(status_code=403, detail="Requiere permisos de administrador de la cuenta")
    return current_user


OrgAdmin = Annotated[Usuario, Depends(require_org_admin)]


def _cuenta_de(db: Session, user: Usuario) -> CuentaCliente:
    if not user.cuenta_id:
        raise HTTPException(status_code=400, detail="Tu usuario no está asociado a una cuenta")
    cuenta = db.get(CuentaCliente, user.cuenta_id)
    if cuenta is None:
        raise HTTPException(status_code=404, detail="Cuenta no encontrada")
    return cuenta


def _empresas_de_cuenta(db: Session, cuenta_id: int) -> list[int]:
    return list(db.scalars(
        select(Empresa.id).where(Empresa.cuenta_id == cuenta_id, Empresa.activa.is_(True))
    ).all())


def _parse_rango(desde: str | None, hasta: str | None) -> tuple[date, date]:
    hoy = date.today()
    try:
        d = date.fromisoformat(desde) if desde else date(hoy.year, 1, 1)
    except ValueError:
        d = date(hoy.year, 1, 1)
    try:
        h = date.fromisoformat(hasta) if hasta else hoy
    except ValueError:
        h = hoy
    if h < d:
        d, h = h, d
    return d, h


@router.get("/proveedores")
def proveedores(
    db: DB,
    admin: OrgAdmin,
    desde: str | None = None,
    hasta: str | None = None,
    empresa_id: int | None = None,
    tarifa: float = 0.0,
    todas_tarifas: bool = False,
    origen: str = "compras",
):
    """Proveedores con operaciones a la tarifa dada, con su concepto
    predominante y los secundarios, cada uno con su clasificación actual.

    `origen` separa la pantalla en dos flujos —"ventas" o "compras"— que nunca
    se mezclan: en ventas exento y excluido siempre quedan separados, en
    compras se agrupan en la presentación principal.

    `todas_tarifas` ignora `tarifa` y trae conceptos de cualquier tarifa —
    hace falta para clasificar IVA descontable (Fase 2), que aplica sobre
    todo a los ítems gravados (5%/19%), no a los de tarifa 0%."""
    if origen not in ("ventas", "compras"):
        raise HTTPException(status_code=400, detail="origen debe ser 'ventas' o 'compras'")

    cuenta = _cuenta_de(db, admin)
    d, h = _parse_rango(desde, hasta)
    empresas = _empresas_de_cuenta(db, cuenta.id)

    if empresa_id is not None:
        if empresa_id not in empresas:
            raise HTTPException(status_code=404, detail="Empresa no encontrada")
        empresas = [empresa_id]

    if not empresas:
        return {"periodo": {"desde": d.isoformat(), "hasta": h.isoformat()}, "proveedores": []}

    resultado = f300.resumen_por_proveedor(
        db, empresa_ids=empresas, desde=d, hasta=h, tarifa=tarifa,
        todas_tarifas=todas_tarifas, origen=origen, cuenta_id=cuenta.id,
    )

    def _concepto(c) -> dict:
        return {
            "concepto": c.concepto,
            "nit_proveedor": c.nit_proveedor,
            "referencia": c.referencia,
            "base_acumulada": c.base_acumulada,
            "documentos": c.documentos,
            "participacion": c.participacion,
            "tratamiento": c.clasificacion.tratamiento,
            "estado": c.clasificacion.estado,
            "origen": c.clasificacion.origen,
            "requiere_revision": c.clasificacion.requiere_revision,
            "es_excepcion": c.clasificacion.es_excepcion,
            "articulo_et": c.clasificacion.articulo_et,
            "norma": c.clasificacion.norma,
            "tipo_item": c.clasificacion.tipo_item,
            "tipo_item_confirmado": c.clasificacion.tipo_item_confirmado,
            "iva_descontable": c.clasificacion.iva_descontable,
            "iva_descontable_confirmado": c.clasificacion.iva_descontable_confirmado,
        }

    return {
        "periodo": {"desde": d.isoformat(), "hasta": h.isoformat()},
        "tarifa": tarifa,
        "todas_tarifas": todas_tarifas,
        "origen": origen,
        "proveedores": [
            {
                "nit": p.nit,
                "razon_social": p.razon_social,
                "base_total": p.base_total,
                "documentos": p.documentos,
                "pendientes": p.pendientes,
                "tipos_documento": p.tipos_documento,
                "predominante": _concepto(p.predominante),
                "secundarios": [_concepto(c) for c in p.secundarios],
            }
            for p in resultado
        ],
    }


@router.get("/reporte")
def reporte(
    db: DB, admin: OrgAdmin,
    desde: str | None = None, hasta: str | None = None,
    empresa_id: int | None = None,
):
    """Resumen consolidado por categoría tributaria para el período dado.
    Incluye ventas (base + IVA generado) y compras (base + IVA facturado +
    IVA descontable), listo para construir la declaración del Formulario 300."""
    cuenta = _cuenta_de(db, admin)
    d, h = _parse_rango(desde, hasta)
    empresas = _empresas_de_cuenta(db, cuenta.id)
    if empresa_id is not None:
        if empresa_id not in empresas:
            raise HTTPException(status_code=404, detail="Empresa no encontrada")
        empresas = [empresa_id]
    return f300.reporte_resumen(db, empresa_ids=empresas, desde=d, hasta=h, cuenta_id=cuenta.id)


@router.get("/tributos-adicionales")
def tributos_adicionales(
    db: DB, admin: OrgAdmin,
    desde: str | None = None, hasta: str | None = None,
    empresa_id: int | None = None,
):
    """Consolida los tributos distintos al IVA (INC, IBUA, ICUI, INPP, bolsas,
    etc.) extraídos de los XML del periodo. Devuelve compras y ventas separadas,
    agrupadas por código DIAN, con drilldown por proveedor."""
    cuenta = _cuenta_de(db, admin)
    d, h = _parse_rango(desde, hasta)
    empresas = _empresas_de_cuenta(db, cuenta.id)
    if empresa_id is not None:
        if empresa_id not in empresas:
            raise HTTPException(status_code=404, detail="Empresa no encontrada")
        empresas = [empresa_id]
    return f300.tributos_adicionales_resumen(
        db, empresa_ids=empresas, desde=d, hasta=h, cuenta_id=cuenta.id
    )


@router.get("/reporte/excel")
def reporte_excel(
    db: DB, admin: OrgAdmin,
    desde: str | None = None, hasta: str | None = None,
    empresa_id: int | None = None,
):
    """Descarga el reporte como archivo Excel (.xlsx)."""
    cuenta = _cuenta_de(db, admin)
    d, h = _parse_rango(desde, hasta)
    empresas = _empresas_de_cuenta(db, cuenta.id)
    nombre_empresa = "Todas las empresas"
    if empresa_id is not None:
        if empresa_id not in empresas:
            raise HTTPException(status_code=404, detail="Empresa no encontrada")
        empresas = [empresa_id]
        emp = db.get(Empresa, empresa_id)
        if emp:
            nombre_empresa = emp.nombre

    nit_empresa = ""
    if empresa_id is not None:
        emp2 = db.get(Empresa, empresa_id)
        if emp2 and emp2.nit:
            nit_empresa = emp2.nit
    datos = f300.reporte_resumen(db, empresa_ids=empresas, desde=d, hasta=h, cuenta_id=cuenta.id)
    tributos = f300.tributos_adicionales_resumen(db, empresa_ids=empresas, desde=d, hasta=h, cuenta_id=cuenta.id)
    buf = _generar_excel(datos, tributos, nombre_empresa, nit_empresa)
    filename = f"formulario300_{d}_{h}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/reporte/pdf")
def reporte_pdf(
    db: DB, admin: OrgAdmin,
    desde: str | None = None, hasta: str | None = None,
    empresa_id: int | None = None,
):
    """Descarga el reporte como archivo PDF."""
    cuenta = _cuenta_de(db, admin)
    d, h = _parse_rango(desde, hasta)
    empresas = _empresas_de_cuenta(db, cuenta.id)
    nombre_empresa = "Todas las empresas"
    if empresa_id is not None:
        if empresa_id not in empresas:
            raise HTTPException(status_code=404, detail="Empresa no encontrada")
        empresas = [empresa_id]
        emp = db.get(Empresa, empresa_id)
        if emp:
            nombre_empresa = emp.nombre

    nit_empresa = ""
    if empresa_id is not None:
        emp2 = db.get(Empresa, empresa_id)
        if emp2 and emp2.nit:
            nit_empresa = emp2.nit
    datos = f300.reporte_resumen(db, empresa_ids=empresas, desde=d, hasta=h, cuenta_id=cuenta.id)
    tributos = f300.tributos_adicionales_resumen(db, empresa_ids=empresas, desde=d, hasta=h, cuenta_id=cuenta.id)
    buf = _generar_pdf(datos, tributos, nombre_empresa, nit_empresa)
    filename = f"formulario300_{d}_{h}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── Generación de exportables ─────────────────────────────────────────────────

def _cop(n: float) -> str:
    return f"${n:,.0f}".replace(",", ".")


def _generar_excel(datos: dict, tributos: dict, nombre_empresa: str, nit_empresa: str = "") -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from datetime import datetime

    # Paleta clara — apta para impresión
    W       = "FFFFFF"; SHEET   = "FFFFFF"
    HDR_BG  = "1E3A5F"; HDR_FG  = "FFFFFF"   # encabezado sección: azul oscuro + blanco
    COL_BG  = "DBEAFE"; COL_FG  = "1E3A5F"   # cabecera columnas: azul pálido + azul oscuro
    ROW_A   = "FFFFFF"; ROW_B   = "F8FAFC"   # filas alternadas blanco / gris muy claro
    BORDE   = "CBD5E1"; DARK    = "1E293B"; MUTED = "64748B"
    VERDE   = "047857"; ROJO_TX = "B91C1C"; DORADO  = "B45309"
    TVB     = "D1FAE5"; TVF     = "065F46"   # total ventas: verde pálido / verde oscuro
    TCB     = "FEE2E2"; TCF     = "991B1B"   # total compras: rojo pálido / rojo oscuro
    TAB     = "FEF3C7"; TAF     = "92400E"   # total adicionales: amber pálido / amber oscuro
    TBALB   = "DBEAFE"; TBALF   = "1E3A5F"   # total balance: azul pálido / azul oscuro
    DEV_BG  = "FEE2E2"                        # fila devoluciones: rojo pálido

    wb = Workbook()
    ws = wb.active
    ws.title = "F300"
    ws.sheet_properties.tabColor = "047857"
    ws.column_dimensions["A"].width = 36
    for col in ["B", "C", "D", "E", "F"]:
        ws.column_dimensions[col].width = 21

    NCOLS = 6
    thin = Side(style="thin", color=BORDE)
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)
    fila = [1]

    def _c(r, c_, val="", bg=W, fg=DARK, bold=False, sz=10, italic=False, ha="left", nf=None):
        cell = ws.cell(row=r, column=c_)
        cell.value = val
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(color=fg, bold=bold, size=sz, italic=italic)
        cell.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=True)
        cell.border = borde
        if nf:
            cell.number_format = nf
        return cell

    def _row_bg(r, bg=W):
        for c_ in range(1, NCOLS + 1):
            cell = ws.cell(row=r, column=c_)
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.border = borde

    def _merged(r, val, bg=W, fg=DARK, bold=False, sz=10, italic=False, ha="center", h=14):
        ws.merge_cells(f"A{r}:F{r}")
        _c(r, 1, val, bg=bg, fg=fg, bold=bold, sz=sz, italic=italic, ha=ha)
        for c_ in range(2, NCOLS + 1):
            ws.cell(row=r, column=c_).fill = PatternFill("solid", fgColor=bg)
        ws.row_dimensions[r].height = h

    def _blank(h=6):
        r = fila[0]; _row_bg(r, bg=W); ws.row_dimensions[r].height = h; fila[0] += 1

    # ── A: Encabezado ──────────────────────────────────────────────────────────
    _merged(fila[0],
            "REPORTE FORMULARIO 300 — CLASIFICACIÓN TRIBUTARIA DE IVA E IMPUESTOS ADICIONALES",
            bg=HDR_BG, fg=HDR_FG, bold=True, sz=12, h=30)
    fila[0] += 1
    _merged(fila[0], f"{nombre_empresa}{f'  ·  NIT {nit_empresa}' if nit_empresa else ''}",
            bg=COL_BG, fg=HDR_BG, sz=10, bold=True, h=18)
    fila[0] += 1
    _merged(fila[0],
            f"Periodo: {datos['periodo']['desde']}  al  {datos['periodo']['hasta']}"
            f"  ·  Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            bg=ROW_B, fg=MUTED, sz=9, h=14)
    fila[0] += 1
    _blank(8)

    # ── B: KPI Cards ───────────────────────────────────────────────────────────
    total_adic = tributos["total_compras"] + tributos["total_ventas"]
    kpis = [
        ("IVA GENERADO", datos["totales"]["iva_generado"], TVB, VERDE),
        ("IVA DESCONTABLE", datos["totales"]["iva_descontable"], TCB, ROJO_TX),
        ("BALANCE ANALÍTICO IVA", datos["totales"]["balance_analitico_iva"], TBALB, TBALF),
        ("TOTAL IMP. ADICIONALES", total_adic, TAB, DORADO),
    ]
    r = fila[0]; ws.row_dimensions[r].height = 14
    for ci, (lbl, _, bg, fg) in enumerate(kpis, 1):
        _c(r, ci, lbl, bg=bg, fg=fg, bold=True, sz=8, ha="center")
    fila[0] += 1
    r = fila[0]; ws.row_dimensions[r].height = 28
    for ci, (_, val, bg, fg) in enumerate(kpis, 1):
        _c(r, ci, val, bg=bg, fg=fg, bold=True, sz=14, ha="center", nf="$ #,##0")
    fila[0] += 1
    _blank(10)

    cats = f300.CATS_LABEL

    def _sec(titulo):
        _merged(fila[0], titulo, bg=HDR_BG, fg=HDR_FG, bold=True, sz=10, ha="left", h=20)
        fila[0] += 1

    def _hdr(*labels):
        r = fila[0]; ws.row_dimensions[r].height = 16
        for ci, lbl in enumerate(labels, 1):
            _c(r, ci, lbl, bg=COL_BG, fg=COL_FG, bold=True, sz=9, ha="right" if ci > 1 else "left")
        for ci in range(len(labels) + 1, NCOLS + 1):
            _c(r, ci, bg=COL_BG)
        fila[0] += 1

    def _dat(*vals, alt=False, fg=DARK, italic=False, bg_ov=None, nf_cols=None):
        r = fila[0]; bg = bg_ov or (ROW_B if alt else ROW_A); ws.row_dimensions[r].height = 14
        for ci, val in enumerate(vals, 1):
            use_nf = "$ #,##0" if (isinstance(val, (int, float)) and ci > 1) else None
            if nf_cols and ci in nf_cols:
                use_nf = nf_cols[ci]
            _c(r, ci, val, bg=bg, fg=fg, italic=italic, sz=9,
               ha="right" if (isinstance(val, (int, float)) and ci > 1) else "left",
               nf=use_nf)
        for ci in range(len(vals) + 1, NCOLS + 1):
            _c(r, ci, bg=bg)
        fila[0] += 1

    def _tot(*vals, bg_t, fg_t):
        r = fila[0]; ws.row_dimensions[r].height = 18
        for ci, val in enumerate(vals, 1):
            _c(r, ci, val, bg=bg_t, fg=fg_t, bold=True, sz=10,
               ha="right" if (isinstance(val, (int, float)) and ci > 1) else "left",
               nf="$ #,##0" if (isinstance(val, (int, float)) and ci > 1) else None)
        for ci in range(len(vals) + 1, NCOLS + 1):
            _c(r, ci, bg=bg_t)
        fila[0] += 1; _blank(8)

    # ── C: Ventas ──────────────────────────────────────────────────────────────
    _sec("MÓDULO 1: VENTAS E INGRESOS — FORMULARIO 300")
    _hdr("Tratamiento / Concepto", "Base gravable (COP)", "IVA generado (COP)")
    v = datos["ventas"]; dev_v = datos.get("devoluciones", {}).get("ventas", {})
    alt = False
    for cat in f300._CATS_REPORTE:
        if v[cat]["base"] == 0 and v[cat]["iva"] == 0:
            continue
        _dat(cats[cat], v[cat]["base"], v[cat]["iva"], alt=alt); alt = not alt
    if dev_v.get("base", 0) > 0:
        _dat("(−) Devoluciones en ventas (NC)", -dev_v["base"], -dev_v.get("iva", 0),
             fg=ROJO_TX, italic=True, bg_ov=DEV_BG)
    _tot("TOTAL NETO", sum(v[c]["base"] for c in f300._CATS_REPORTE),
         datos["totales"]["iva_generado"], bg_t=TVB, fg_t=TVF)

    # ── D: Compras ─────────────────────────────────────────────────────────────
    _sec("MÓDULO 2: COMPRAS Y COSTOS — FORMULARIO 300")
    _hdr("Tratamiento / Concepto", "Base (COP)", "IVA facturado (COP)",
         "IVA descontable (COP)", "Mayor valor costo/gasto (COP)")
    c = datos["compras"]; dev_c = datos.get("devoluciones", {}).get("compras", {})
    alt = False
    for cat in f300._CATS_REPORTE:
        if c[cat]["base"] == 0 and c[cat]["iva_facturado"] == 0:
            continue
        mayor = c[cat]["iva_facturado"] - c[cat]["iva_descontable"]
        _dat(cats[cat], c[cat]["base"], c[cat]["iva_facturado"],
             c[cat]["iva_descontable"], mayor, alt=alt); alt = not alt
    if dev_c.get("base", 0) > 0:
        dv_mayor = dev_c.get("iva_facturado", 0) - dev_c.get("iva_descontable", 0)
        _dat("(−) Devoluciones en compras (NC)", -dev_c["base"],
             -dev_c.get("iva_facturado", 0), -dev_c.get("iva_descontable", 0), -dv_mayor,
             fg=ROJO_TX, italic=True, bg_ov=DEV_BG)
    _tot("TOTAL NETO",
         sum(c[cat]["base"] for cat in f300._CATS_REPORTE),
         sum(c[cat]["iva_facturado"] for cat in f300._CATS_REPORTE),
         datos["totales"]["iva_descontable"],
         sum(c[cat]["iva_facturado"] - c[cat]["iva_descontable"] for cat in f300._CATS_REPORTE),
         bg_t=TCB, fg_t=TCF)

    # ── E: Impuestos adicionales ───────────────────────────────────────────────
    _sec("MÓDULO 3: IMPUESTOS ADICIONALES Y TRIBUTOS SALUDABLES")
    todos: dict[str, dict] = {}
    for t in tributos.get("compras", []):
        todos.setdefault(t["cod_dian"], {"nombre": t["nombre"], "compras": 0.0, "ventas": 0.0})
        todos[t["cod_dian"]]["compras"] += t["valor_total"]
    for t in tributos.get("ventas", []):
        todos.setdefault(t["cod_dian"], {"nombre": t["nombre"], "compras": 0.0, "ventas": 0.0})
        todos[t["cod_dian"]]["ventas"] += t["valor_total"]
    if todos:
        _hdr("Tributo", "Código DIAN", "Compras (COP)", "Ventas (COP)", "Total (COP)")
        alt = False
        for cod, entry in sorted(todos.items(), key=lambda e: e[1]["compras"] + e[1]["ventas"], reverse=True):
            _dat(entry["nombre"], cod, entry["compras"], entry["ventas"],
                 entry["compras"] + entry["ventas"], alt=alt)
            alt = not alt
        _tot("TOTAL TRIBUTOS ADICIONALES", "",
             tributos["total_compras"], tributos["total_ventas"],
             tributos["total_compras"] + tributos["total_ventas"],
             bg_t=TAB, fg_t=TAF)
    else:
        _merged(fila[0], "No se encontraron impuestos adicionales en el periodo.",
                bg=ROW_B, fg=MUTED, sz=9, ha="left", h=15)
        fila[0] += 1; _blank(8)

    # ── F: Balance + notas ─────────────────────────────────────────────────────
    _sec("MÓDULO 4: BALANCE ANALÍTICO DE IVA")
    _dat("IVA generado (ventas)", datos["totales"]["iva_generado"], fg=VERDE)
    _dat("IVA descontable (compras)", datos["totales"]["iva_descontable"], fg=ROJO_TX, alt=True)
    _tot("Balance analítico de IVA", datos["totales"]["balance_analitico_iva"],
         bg_t=TBALB, fg_t=TBALF)

    notas = [
        "* Este valor no equivale al saldo a pagar ni al saldo a favor final. El Formulario 300 "
        "depende de otros conceptos de la liquidación (saldos a favor previos, retenciones "
        "practicadas, prorrateos del Art. 490 ET) que este análisis de XML no cubre.",
        "* Los tributos adicionales están discriminados directamente de los XML. Ciolix los muestra "
        "para que el contador determine su tratamiento fiscal: mayor valor del costo/gasto, "
        "impuesto recuperable u otro según la norma aplicable.",
    ]
    for nota in notas:
        r = fila[0]
        ws.merge_cells(f"A{r}:F{r}")
        cell = ws.cell(row=r, column=1)
        cell.value = nota
        cell.fill = PatternFill("solid", fgColor=ROW_B)
        cell.font = Font(color=MUTED, size=8, italic=True)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = borde
        ws.row_dimensions[r].height = 32
        for c_ in range(2, NCOLS + 1):
            ws.cell(row=r, column=c_).fill = PatternFill("solid", fgColor=ROW_B)
        fila[0] += 1

    if datos.get("conceptos_pendientes", 0) > 0:
        _merged(fila[0],
                f"⚠ {datos['conceptos_pendientes']} concepto(s) sin clasificar — "
                "las cifras de exento/excluido pueden estar incompletas.",
                bg=TAB, fg=TAF, sz=9, ha="left", h=16)
        fila[0] += 1
    if tributos.get("hay_no_parametrizados"):
        _merged(fila[0], "⚠ Hay tributos con código DIAN no parametrizado — requieren revisión manual.",
                bg=TAB, fg=TAF, sz=9, ha="left", h=16)
        fila[0] += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _generar_pdf(datos: dict, tributos: dict, nombre_empresa: str, nit_empresa: str = "") -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from datetime import datetime

    # Paleta clara — apta para impresión
    W      = colors.white
    HDR_BG = colors.HexColor("#1E3A5F"); HDR_FG = colors.white
    COL_BG = colors.HexColor("#DBEAFE"); COL_FG = colors.HexColor("#1E3A5F")
    ROW_A  = colors.white;               ROW_B  = colors.HexColor("#F8FAFC")
    BORDE  = colors.HexColor("#CBD5E1"); DARK   = colors.HexColor("#1E293B")
    MUTED  = colors.HexColor("#64748B")
    VERDE  = colors.HexColor("#047857"); ROJO   = colors.HexColor("#B91C1C")
    DORADO = colors.HexColor("#B45309")
    TVB    = colors.HexColor("#D1FAE5"); TVF    = colors.HexColor("#065F46")
    TCB    = colors.HexColor("#FEE2E2"); TCF    = colors.HexColor("#991B1B")
    TAB    = colors.HexColor("#FEF3C7"); TAF    = colors.HexColor("#92400E")
    TBALB  = colors.HexColor("#DBEAFE"); TBALF  = colors.HexColor("#1E3A5F")
    DEV_BG = colors.HexColor("#FEE2E2")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.2*cm, rightMargin=1.2*cm,
                            topMargin=1.2*cm, bottomMargin=1.2*cm)

    def _p(text, color=DARK, sz=9, bold=False, italic=False, align=0):
        fn = "Helvetica-Bold" if bold else ("Helvetica-Oblique" if italic else "Helvetica")
        return Paragraph(text, ParagraphStyle("x", fontName=fn, fontSize=sz,
                                              textColor=color, leading=sz + 2, alignment=align))

    def _th(titulo, w=None):
        t = Table([[_p(titulo, HDR_FG, 10, bold=True)]], colWidths=[w or doc.width])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), HDR_BG),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        return t

    def _tabla(hdrs, rows, total_row, col_w, total_bg, total_fg, dev_idx=None):
        data = [hdrs] + rows + [total_row]
        ts = TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), COL_BG),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 0), (0, -1), "LEFT"),
            ("BACKGROUND", (0, -1), (-1, -1), total_bg),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("TEXTCOLOR", (0, -1), (-1, -1), total_fg),
            ("GRID", (0, 0), (-1, -1), 0.4, BORDE),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (1, 0), (-1, -1), 5),
        ])
        for i in range(1, len(rows) + 1):
            ts.add("BACKGROUND", (0, i), (-1, i), ROW_B if i % 2 == 0 else ROW_A)
        if dev_idx is not None:
            ts.add("BACKGROUND", (0, dev_idx), (-1, dev_idx), DEV_BG)
        t = Table(data, colWidths=col_w)
        t.setStyle(ts)
        return t

    cats = f300.CATS_LABEL
    story = []

    # ── A: Encabezado ──────────────────────────────────────────────────────────
    empresa_str = f"{nombre_empresa}{f'  ·  NIT {nit_empresa}' if nit_empresa else ''}"
    periodo_str = (f"Periodo: {datos['periodo']['desde']}  al  {datos['periodo']['hasta']}"
                   f"  ·  Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    tbl_enc = Table([
        [_p("REPORTE FORMULARIO 300 — CLASIFICACIÓN TRIBUTARIA DE IVA E IMPUESTOS ADICIONALES",
            HDR_FG, 12, bold=True, align=1)],
        [_p(empresa_str, COL_FG, 10, bold=True, align=1)],
        [_p(periodo_str, MUTED, 8, align=1)],
    ], colWidths=[doc.width])
    tbl_enc.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), HDR_BG),
        ("BACKGROUND", (0, 1), (0, 1), COL_BG),
        ("BACKGROUND", (0, 2), (0, 2), ROW_B),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDE),
    ]))
    story.extend([tbl_enc, Spacer(1, 0.4*cm)])

    # ── B: KPIs ────────────────────────────────────────────────────────────────
    total_adic = tributos["total_compras"] + tributos["total_ventas"]
    kpis = [
        ("IVA GENERADO", datos["totales"]["iva_generado"], TVB, TVF),
        ("IVA DESCONTABLE", datos["totales"]["iva_descontable"], TCB, TCF),
        ("BALANCE ANALÍTICO IVA", datos["totales"]["balance_analitico_iva"], TBALB, TBALF),
        ("IMP. ADICIONALES", total_adic, TAB, TAF),
    ]
    kw = doc.width / 4
    tbl_kpi = Table(
        [[_p(k[0], k[3], 7, bold=True, align=1) for k in kpis],
         [_p(_cop(k[1]), k[3], 13, bold=True, align=1) for k in kpis]],
        colWidths=[kw] * 4,
    )
    kpi_style = [("ALIGN", (0, 0), (-1, -1), "CENTER"),
                 ("GRID", (0, 0), (-1, -1), 0.4, BORDE),
                 ("TOPPADDING", (0, 0), (-1, 0), 6), ("BOTTOMPADDING", (0, 0), (-1, 0), 2),
                 ("TOPPADDING", (0, 1), (-1, 1), 2), ("BOTTOMPADDING", (0, 1), (-1, 1), 8)]
    for i, k in enumerate(kpis):
        kpi_style.append(("BACKGROUND", (i, 0), (i, 1), k[2]))
    tbl_kpi.setStyle(TableStyle(kpi_style))
    story.extend([tbl_kpi, Spacer(1, 0.4*cm)])

    # ── C: Ventas ──────────────────────────────────────────────────────────────
    v = datos["ventas"]; dev_v = datos.get("devoluciones", {}).get("ventas", {})
    rows_v = [
        [_p(cats[c], DARK, 8), _p(_cop(v[c]["base"]), DARK, 8), _p(_cop(v[c]["iva"]), VERDE, 8)]
        for c in f300._CATS_REPORTE if v[c]["base"] != 0 or v[c]["iva"] != 0
    ]
    dev_idx_v = None
    if dev_v.get("base", 0) > 0:
        dev_idx_v = len(rows_v) + 1
        rows_v.append([_p("(−) Devoluciones en ventas (NC)", ROJO, 8, italic=True),
                       _p(f'−{_cop(dev_v["base"])}', ROJO, 8, italic=True),
                       _p(f'−{_cop(dev_v.get("iva", 0))}', ROJO, 8, italic=True)])
    story.append(_th("MÓDULO 1: VENTAS E INGRESOS — FORMULARIO 300"))
    story.append(_tabla(
        hdrs=[_p("Tratamiento / Concepto", COL_FG, 8, bold=True),
              _p("Base gravable (COP)", COL_FG, 8, bold=True),
              _p("IVA generado (COP)", COL_FG, 8, bold=True)],
        rows=rows_v,
        total_row=[_p("TOTAL NETO", TVF, 9, bold=True),
                   _p(_cop(sum(v[c]["base"] for c in f300._CATS_REPORTE)), TVF, 9, bold=True),
                   _p(_cop(datos["totales"]["iva_generado"]), TVF, 9, bold=True)],
        col_w=[10*cm, 5*cm, 5*cm],
        total_bg=TVB, total_fg=TVF, dev_idx=dev_idx_v,
    ))
    story.append(Spacer(1, 0.3*cm))

    # ── D: Compras ─────────────────────────────────────────────────────────────
    c = datos["compras"]; dev_c = datos.get("devoluciones", {}).get("compras", {})
    rows_c = []
    for cat in f300._CATS_REPORTE:
        if c[cat]["base"] == 0 and c[cat]["iva_facturado"] == 0:
            continue
        mayor = c[cat]["iva_facturado"] - c[cat]["iva_descontable"]
        rows_c.append([_p(cats[cat], DARK, 8),
                       _p(_cop(c[cat]["base"]), DARK, 8),
                       _p(_cop(c[cat]["iva_facturado"]), DARK, 8),
                       _p(_cop(c[cat]["iva_descontable"]), VERDE, 8),
                       _p(_cop(mayor), ROJO, 8)])
    dev_idx_c = None
    if dev_c.get("base", 0) > 0:
        dev_idx_c = len(rows_c) + 1
        dv_mayor = dev_c.get("iva_facturado", 0) - dev_c.get("iva_descontable", 0)
        rows_c.append([_p("(−) Devoluciones en compras (NC)", ROJO, 8, italic=True),
                       _p(f'−{_cop(dev_c["base"])}', ROJO, 8, italic=True),
                       _p(f'−{_cop(dev_c.get("iva_facturado", 0))}', ROJO, 8, italic=True),
                       _p(f'−{_cop(dev_c.get("iva_descontable", 0))}', ROJO, 8, italic=True),
                       _p(f'−{_cop(dv_mayor)}', ROJO, 8, italic=True)])
    story.append(_th("MÓDULO 2: COMPRAS Y COSTOS — FORMULARIO 300"))
    story.append(_tabla(
        hdrs=[_p("Tratamiento / Concepto", COL_FG, 8, bold=True),
              _p("Base (COP)", COL_FG, 8, bold=True),
              _p("IVA facturado (COP)", COL_FG, 8, bold=True),
              _p("IVA descontable (COP)", COL_FG, 8, bold=True),
              _p("Mayor valor costo/gasto (COP)", COL_FG, 8, bold=True)],
        rows=rows_c,
        total_row=[_p("TOTAL NETO", TCF, 9, bold=True),
                   _p(_cop(sum(c[cat]["base"] for cat in f300._CATS_REPORTE)), TCF, 9, bold=True),
                   _p(_cop(sum(c[cat]["iva_facturado"] for cat in f300._CATS_REPORTE)), TCF, 9, bold=True),
                   _p(_cop(datos["totales"]["iva_descontable"]), TVF, 9, bold=True),
                   _p(_cop(sum(c[cat]["iva_facturado"] - c[cat]["iva_descontable"]
                               for cat in f300._CATS_REPORTE)), TCF, 9, bold=True)],
        col_w=[8*cm, 3.8*cm, 3.8*cm, 3.8*cm, 4.1*cm],
        total_bg=TCB, total_fg=TCF, dev_idx=dev_idx_c,
    ))
    story.append(Spacer(1, 0.3*cm))

    # ── E: Impuestos adicionales ───────────────────────────────────────────────
    todos: dict[str, dict] = {}
    for t in tributos.get("compras", []):
        todos.setdefault(t["cod_dian"], {"nombre": t["nombre"], "compras": 0.0, "ventas": 0.0})
        todos[t["cod_dian"]]["compras"] += t["valor_total"]
    for t in tributos.get("ventas", []):
        todos.setdefault(t["cod_dian"], {"nombre": t["nombre"], "compras": 0.0, "ventas": 0.0})
        todos[t["cod_dian"]]["ventas"] += t["valor_total"]

    story.append(_th("MÓDULO 3: IMPUESTOS ADICIONALES Y TRIBUTOS SALUDABLES"))
    if todos:
        rows_adic = [
            [_p(entry["nombre"], DARK, 8), _p(cod, MUTED, 8),
             _p(_cop(entry["compras"]), ROJO, 8), _p(_cop(entry["ventas"]), VERDE, 8),
             _p(_cop(entry["compras"] + entry["ventas"]), DORADO, 8)]
            for cod, entry in sorted(todos.items(), key=lambda e: e[1]["compras"] + e[1]["ventas"], reverse=True)
        ]
        story.append(_tabla(
            hdrs=[_p("Tributo", COL_FG, 8, bold=True), _p("Código", COL_FG, 8, bold=True),
                  _p("Compras (COP)", COL_FG, 8, bold=True), _p("Ventas (COP)", COL_FG, 8, bold=True),
                  _p("Total (COP)", COL_FG, 8, bold=True)],
            rows=rows_adic,
            total_row=[_p("TOTAL TRIBUTOS ADICIONALES", TAF, 9, bold=True), _p("", TAF, 9),
                       _p(_cop(tributos["total_compras"]), TAF, 9, bold=True),
                       _p(_cop(tributos["total_ventas"]), TAF, 9, bold=True),
                       _p(_cop(tributos["total_compras"] + tributos["total_ventas"]), TAF, 9, bold=True)],
            col_w=[9*cm, 2.5*cm, 4.5*cm, 4.5*cm, 4.5*cm],
            total_bg=TAB, total_fg=TAF,
        ))
    else:
        story.append(_p("No se encontraron impuestos adicionales en el periodo.", MUTED, 9))
    story.append(Spacer(1, 0.3*cm))

    # ── F: Balance + notas ─────────────────────────────────────────────────────
    story.append(_th("MÓDULO 4: BALANCE ANALÍTICO DE IVA"))
    tbl_bal = Table([
        [_p("IVA generado (ventas)", VERDE, 9), _p(_cop(datos["totales"]["iva_generado"]), VERDE, 9, align=2)],
        [_p("IVA descontable (compras)", ROJO, 9), _p(_cop(datos["totales"]["iva_descontable"]), ROJO, 9, align=2)],
        [_p("Balance analítico de IVA", TBALF, 10, bold=True),
         _p(_cop(datos["totales"]["balance_analitico_iva"]), TBALF, 10, bold=True, align=2)],
    ], colWidths=[9*cm, 5*cm])
    tbl_bal.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ROW_A),
        ("BACKGROUND", (0, 1), (-1, 1), ROW_B),
        ("BACKGROUND", (0, 2), (-1, 2), TBALB),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDE),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (-1, 0), (-1, -1), 8),
    ]))
    story.extend([tbl_bal, Spacer(1, 0.3*cm)])

    nota_st = ParagraphStyle("n", fontName="Helvetica-Oblique", fontSize=8,
                             textColor=MUTED, leading=10, spaceAfter=4)
    story.append(Paragraph(
        "* Este valor no equivale al saldo a pagar ni al saldo a favor final. El Formulario 300 "
        "depende de otros conceptos de la liquidación (saldos a favor previos, retenciones "
        "practicadas, prorrateos del Art. 490 ET) que este análisis de XML no cubre.", nota_st))
    story.append(Paragraph(
        "* Los tributos adicionales están discriminados directamente de los XML. Ciolix los muestra "
        "para que el contador determine su tratamiento fiscal: mayor valor del costo/gasto, "
        "impuesto recuperable u otro según la norma aplicable.", nota_st))

    warn_st = ParagraphStyle("w", fontName="Helvetica-Oblique", fontSize=8,
                             textColor=DORADO, leading=10)
    if datos.get("conceptos_pendientes", 0) > 0:
        story.append(Paragraph(
            f"⚠ {datos['conceptos_pendientes']} concepto(s) sin clasificar — "
            "las cifras de exento/excluido pueden estar incompletas.", warn_st))
    if tributos.get("hay_no_parametrizados"):
        story.append(Paragraph(
            "⚠ Hay tributos con código DIAN no parametrizado — requieren revisión manual.", warn_st))

    doc.build(story)
    buf.seek(0)
    return buf


class ClasificarRequest(BaseModel):
    empresa_id: int
    concepto: str = Field(..., min_length=1)
    tratamiento: str
    # Con NIT: la decisión solo vale para ese proveedor (queda como excepción,
    # no se comparte). Sin NIT: es una regla general del concepto y se propone
    # también a las demás empresas de la firma.
    nit_tercero: str | None = None
    # Código/referencia del producto (si el XML la trae) — afina la memoria
    # cuando el mismo proveedor vende bienes y servicios con textos parecidos.
    referencia: str | None = None
    # bien | servicio, confirmado por el contador — opcional, es un eje
    # independiente del tratamiento de IVA.
    tipo_item: str | None = None
    # descontable | no_descontable, confirmado por el contador — opcional,
    # Fase 2. Nunca se propaga al catálogo de la firma (ver
    # validar_clasificacion).
    iva_descontable: str | None = None
    articulo_et: str | None = None
    norma: str | None = None


@router.post("/clasificar")
def clasificar_endpoint(body: ClasificarRequest, db: DB, admin: OrgAdmin):
    cuenta = _cuenta_de(db, admin)
    empresa = db.get(Empresa, body.empresa_id)
    if empresa is None or empresa.cuenta_id != cuenta.id:
        raise HTTPException(status_code=404, detail="Empresa no encontrada")

    if body.tratamiento not in TRATAMIENTOS:
        raise HTTPException(
            status_code=400,
            detail=f"Tratamiento inválido. Debe ser uno de: {', '.join(TRATAMIENTOS)}",
        )
    if body.tipo_item is not None and body.tipo_item not in TIPOS_ITEM:
        raise HTTPException(
            status_code=400,
            detail=f"Tipo de ítem inválido. Debe ser uno de: {', '.join(TIPOS_ITEM)}",
        )
    if body.iva_descontable is not None and body.iva_descontable not in IVA_DESCONTABLE_ESTADOS:
        raise HTTPException(
            status_code=400,
            detail=f"IVA descontable inválido. Debe ser uno de: {', '.join(IVA_DESCONTABLE_ESTADOS)}",
        )

    fila = f300.validar_clasificacion(
        db, empresa=empresa, usuario=admin, concepto=body.concepto,
        tratamiento=body.tratamiento, nit_tercero=body.nit_tercero,
        referencia=body.referencia, tipo_item=body.tipo_item,
        iva_descontable=body.iva_descontable,
        articulo_et=body.articulo_et, norma=body.norma,
    )
    return {
        "id": fila.id, "concepto": fila.concepto, "tratamiento": fila.tratamiento,
        "estado": fila.estado, "es_excepcion": fila.es_excepcion,
        "tipo_item": fila.tipo_item, "iva_descontable": fila.iva_descontable,
    }
